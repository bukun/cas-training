from openpyxl import load_workbook

excel = load_workbook('./demo_tmpl/e_mod.xlsx')
ws = excel.active
sheet = excel.get_sheet_by_name('Sheet1')

print(sheet.max_row)
print(sheet.max_column)

for row in sheet.rows:
    # print(row)
    pass

array_20200420 = []
for row_idx in range(4, sheet.max_row + 1):
    # print(sheet.cell(row = row_idx, column = 1).value)
    array_20200420.append(sheet.cell(row = row_idx, column = 2).value)

import math
avg = sum(array_20200420) / len(array_20200420)
print(avg)

sheet.cell(row = 1, column = 10).value = avg

excel.save('xx_data.xlsx')